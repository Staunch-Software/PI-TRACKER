"""Excel export helper for the Landing Reports module (see api/routes/landing_reports.py)."""
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXPORT_HEADERS = ["S.No", "Vessel Name", "Port", "Oil Type", "Offlanded Date", "Shipment Status", "Report Status"]

_HEADER_FILL = PatternFill("solid", fgColor="0F4C81")
_HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)


def build_export_workbook(rows: list[tuple]) -> BytesIO:
    """rows: (LandingReport, vessel_name) pairs, already filtered and sorted."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Landing Reports"
    sheet.append(EXPORT_HEADERS)
    for col_idx in range(1, len(EXPORT_HEADERS) + 1):
        cell = sheet.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(col_idx)].width = 22
    for report, vessel_name in rows:
        sheet.append(
            [
                report.seq_no,
                vessel_name,
                report.port,
                report.oil_type,
                report.offlanded_date,
                report.shipment_status,
                report.report_status,
            ]
        )
    for row in sheet.iter_rows(min_row=2, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = "DD-MM-YYYY"
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
