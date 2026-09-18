import re
from io import BytesIO

import openpyxl
from sqlalchemy.orm import Session

from app.core.enums import Department
from app.models.vendor_department_mapping import VendorDepartmentMapping
from app.schemas.vendor_department_mapping import ColumnMappingCandidate, VendorImportRowPreview

# Header text (normalized: lowercased, punctuation stripped to spaces, whitespace collapsed) ->
# field name. Real-world sheets vary a lot in header wording, so this is intentionally generous;
# anything not confidently matched falls through to the column-mapping confirmation step instead
# of being silently guessed.
_VENDOR_NAME_HEADERS = {"vendor name", "vendor", "vendor_name", "vendorname", "supplier", "supplier name", "name"}
_DEPARTMENT_HEADERS = {"department", "dept", "vendor department", "division", "category", "vendor type", "type"}

# Accepted spellings for each Department value (normalized: lowercased, trimmed).
_DEPARTMENT_VALUES: dict[str, Department] = {
    "technical": Department.TECHNICAL,
    "tech": Department.TECHNICAL,
    "manning": Department.MANNING,
    "crew": Department.MANNING,
    "crewing": Department.MANNING,
}


def normalize_header(raw: str) -> str:
    lowered = re.sub(r"[^a-z0-9\s]", " ", raw.lower())
    return re.sub(r"\s+", " ", lowered).strip()


def normalize_vendor_name(raw: str) -> str:
    """Case/whitespace/punctuation-insensitive key used for duplicate detection and upsert."""
    return re.sub(r"[^a-z0-9]", "", raw.lower())


def parse_department(value: object) -> tuple[Department | None, str | None]:
    if value is None or str(value).strip() == "":
        return None, "department is required"
    normalized = re.sub(r"\s+", " ", str(value).strip().lower())
    department = _DEPARTMENT_VALUES.get(normalized)
    if department is None:
        return None, f"department '{value}' not recognized (expected Technical or Manning)"
    return department, None


def _guess_field(normalized_header: str) -> str | None:
    if normalized_header in _VENDOR_NAME_HEADERS:
        return "vendor_name"
    if normalized_header in _DEPARTMENT_HEADERS:
        return "department"
    return None


def read_headers(file_bytes: bytes) -> tuple[list[ColumnMappingCandidate], object]:
    """Returns the first row's headers (with best-guess field per column) and the worksheet."""
    workbook = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    sheet = workbook.active

    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers: list[ColumnMappingCandidate] = []
    for col_idx, raw_header in enumerate(header_row):
        if raw_header is None or str(raw_header).strip() == "":
            continue
        headers.append(
            ColumnMappingCandidate(
                raw_header=str(raw_header).strip(),
                column_index=col_idx,
                guessed_field=_guess_field(normalize_header(str(raw_header))),
            )
        )
    return headers, sheet


def resolve_confident_mapping(headers: list[ColumnMappingCandidate]) -> dict[str, int] | None:
    """If exactly one column guesses vendor_name and exactly one guesses department, we're
    confident enough to proceed without asking the user. Otherwise return None so the caller
    can show a column-mapping confirmation step instead of guessing wrong."""
    vendor_cols = [h.column_index for h in headers if h.guessed_field == "vendor_name"]
    department_cols = [h.column_index for h in headers if h.guessed_field == "department"]
    if len(vendor_cols) == 1 and len(department_cols) == 1:
        return {"vendor_name": vendor_cols[0], "department": department_cols[0]}
    return None


def parse_rows(
    db: Session, sheet, column_mapping: dict[str, int]
) -> list[VendorImportRowPreview]:
    vendor_col = column_mapping["vendor_name"]
    department_col = column_mapping["department"]

    existing_normalized = {n for (n,) in db.query(VendorDepartmentMapping.vendor_name_normalized).all()}

    previews: list[VendorImportRowPreview] = []
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        raw_name = row[vendor_col] if vendor_col < len(row) else None
        raw_department = row[department_col] if department_col < len(row) else None

        errors: list[str] = []
        vendor_name = str(raw_name).strip() if raw_name is not None else ""
        if not vendor_name:
            errors.append("vendor name is required")

        department, err = parse_department(raw_department)
        if err:
            errors.append(err)

        is_update = bool(vendor_name) and normalize_vendor_name(vendor_name) in existing_normalized

        previews.append(
            VendorImportRowPreview(
                row_number=row_idx,
                vendor_name=vendor_name or None,
                department=department,
                errors=errors,
                is_update=is_update,
            )
        )

    return previews
